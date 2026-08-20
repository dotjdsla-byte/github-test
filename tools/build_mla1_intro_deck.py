# -*- coding: utf-8 -*-
"""MLA Level 1 — part 1 (fundamentals), in the KLC-SM house style.

Built from the Lube_Win_MLA_1_1st notes and the screen captures embedded in
them. Four slides: strategy, friction regimes, base oils, additives/viscosity/
grease. Body text is English so it can be reviewed with an English-speaking
manager; the captures are reused as-is.
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import deck_style as ds
from deck_style import (L, R, COLW, TOP, INK, GREY, ACCENT, NAVY, WHITE, LIGHT,
                        chip, textbox, para, rect, table, callout, blank_slide)

ds.LATIN_FONT = "Calibri"

MEDIA = "/tmp/claude-0/-home-user-github-test/eae68b7d-8c95-56e3-85c5-ff5decc49086/scratchpad/mla1/"
OUT = "/home/user/github-test/MLA1_Fundamentals_Part1.pptx"

prs = Presentation()
prs.slide_width, prs.slide_height = ds.SW, ds.SH

# ---------------------------------------------------------------- slide 1
s = blank_slide(prs, "Lubrication Strategy Follows Criticality")

y = chip(s, L, TOP, COLW, "Why lubrication is preventive maintenance")
tb, tf = textbox(s, L, y, COLW, Inches(1.5))
para(tf, "If roughly 80% of failures trace back to contaminated or degraded oil, "
         "then every lubrication action counts as preventive maintenance — not housekeeping.",
     13, False, INK, 8, first=True, line=1.25)
para(tf, "The same logic as heart health: diet and a blood test every six months "
         "cost very little. Surgery does not.", 13, False, GREY, 0, line=1.25)

y2 = y + Inches(1.55)
y2 = chip(s, L, y2, COLW, "The six functions of a lubricant")
table(s, L, y2, COLW,
      [[("Reduce friction", False, INK, PP_ALIGN.CENTER), ("Reduce wear", False, INK, PP_ALIGN.CENTER)],
       [("Prevent corrosion", False, INK, PP_ALIGN.CENTER), ("Remove heat", False, INK, PP_ALIGN.CENTER)],
       [("Remove contaminants", False, INK, PP_ALIGN.CENTER), ("Transmit fluid power", False, INK, PP_ALIGN.CENTER)]],
      [Inches(2.975), Inches(2.975)], head=False, row_h=Inches(0.42), size=12.5)

y = chip(s, R, TOP, COLW, "There is no single right answer — spend where it counts")
end = table(s, R, y, COLW,
            [["Equipment character", "What to specify"],
             [("Cheap, easy to replace,\nnot critical", False, INK, PP_ALIGN.LEFT),
              ("A basic lubricant is enough", False, INK, PP_ALIGN.LEFT)],
             [("Spare held on board,\nquick swap", False, INK, PP_ALIGN.LEFT),
              ("Standard grade", False, INK, PP_ALIGN.LEFT)],
             [("Stern tube, thruster —\ncostly, hard to reach,\nexpensive downtime", True, ACCENT, PP_ALIGN.LEFT),
              ("Specify the best available", True, ACCENT, PP_ALIGN.LEFT)]],
            [Inches(2.85), Inches(3.10)], row_h=Inches(0.62), size=12)

callout(s, R, end + Inches(0.25), COLW,
        "Criticality — not price — sets the specification.", h=Inches(0.7),
        size=14, bold=True, color=NAVY)

# ---------------------------------------------------------------- slide 2
s = blank_slide(prs, "Friction and the Four Lubrication Regimes")

y = chip(s, L, TOP, COLW, "Stribeck curve")
s.shapes.add_picture(MEDIA + "image5.png", L, y, width=COLW)   # 1264x634 -> 2.98"
tb, tf = textbox(s, L, y + Inches(3.12), COLW, Inches(0.32))
para(tf, "Friction against ZN/P, and where each regime shows up on board",
     12, False, GREY, 0, first=True, align=PP_ALIGN.CENTER)

callout(s, L, y + Inches(3.60), COLW,
        "Slow and heavily loaded \u2192 higher viscosity, so the film forms sooner.",
        h=Inches(0.7), size=14, bold=True, color=NAVY)

y = chip(s, R, TOP, COLW, "The intuition")
s.shapes.add_picture(MEDIA + "image3.png", R, y, width=COLW)   # 1386x673 -> 2.89"
tb, tf = textbox(s, R, y + Inches(3.02), COLW, Inches(1.5))
para(tf, "Push a block across concrete and friction is highest before it moves — "
         "and barely depends on speed.", 13, False, INK, 8, first=True, line=1.25)
para(tf, "Add a lubricant and friction drops sharply. Keep raising the speed and it "
         "climbs again, the way a ship has to push aside more water the faster it goes: "
         "now the oil itself is what must be sheared.", 13, False, INK, 0, line=1.25)


# ---------------------------------------------------------------- slide 3
s = blank_slide(prs, "Base Oils — Mineral and Synthetic")

y = chip(s, L, TOP, COLW, "Two ways to make a base oil")
s.shapes.add_picture(MEDIA + "image9_crop.png", L, y, width=COLW)  # 1320x583 -> 2.63"
tb, tf = textbox(s, L, y + Inches(2.78), COLW, Inches(1.5))
para(tf, "Mineral — refined from crude oil. A box of mixed Lego: many shapes, "
         "many sizes, whatever the barrel gave you.", 13, False, INK, 8, first=True, line=1.25)
para(tf, "Synthetic — known molecules assembled into the structure you actually want.",
     13, False, INK, 0, line=1.25)

y = chip(s, R, TOP, COLW, "API base oil groups")
s.shapes.add_picture(MEDIA + "image10.png", R, y, width=COLW)  # 1301x632 -> 2.89"

yy = y + Inches(3.05)
tb, tf = textbox(s, R, yy, COLW, Inches(1.6))
para(tf, "Group III is sold as synthetic in most markets, but it is still refined "
         "mineral stock — synthetic in name.", 12.5, False, INK, 6, first=True, line=1.2)
para(tf, "Group IV (PAO) is the workhorse synthetic — gears and some compressors.",
     12.5, False, INK, 6, line=1.2)
para(tf, "Group V is everything that fits nowhere else, which does not make it better.",
     12.5, False, INK, 0, line=1.2)

callout(s, R, yy + Inches(1.20), COLW,
        "Viscosity Index = how much the oil thins as it heats. Higher VI, flatter response.",
        h=Inches(0.65), size=13, bold=True, color=NAVY)

# ---------------------------------------------------------------- slide 4
s = blank_slide(prs, "Additives, Viscosity and Grease")

y = chip(s, L, TOP, COLW, "Additive families")
s.shapes.add_picture(MEDIA + "image11_crop.png", L + Inches(0.62), y,
                     width=Inches(4.70))                  # 1009x641 -> 2.99"

y2 = y + Inches(3.12)
tb, tf = textbox(s, L, y2, COLW, Inches(0.55))
para(tf, "Additives are the most expensive part of the blend — do not pay for what "
         "the application does not need.", 12.5, True, ACCENT, 0, first=True, line=1.2)

table(s, L, y2 + Inches(0.58), COLW,
      [[("Open to atmosphere", False, INK, PP_ALIGN.LEFT), ("More oxidation inhibitor", False, INK, PP_ALIGN.LEFT)],
       [("Heavily loaded", False, INK, PP_ALIGN.LEFT), ("Extreme pressure (EP)", False, INK, PP_ALIGN.LEFT)],
       [("Hydraulic pump", False, INK, PP_ALIGN.LEFT), ("Antiwear (AW)", False, INK, PP_ALIGN.LEFT)],
       [("Compressor, turbine", False, INK, PP_ALIGN.LEFT), ("R&O — rust and oxidation", False, INK, PP_ALIGN.LEFT)]],
      [Inches(2.45), Inches(3.50)], head=False, row_h=Inches(0.32), size=12)

y = chip(s, R, TOP, COLW, "The three properties that matter most")
end = table(s, R, y, COLW,
            [[("1", True, NAVY, PP_ALIGN.CENTER), ("Viscosity", False, INK, PP_ALIGN.LEFT)],
             [("2", True, NAVY, PP_ALIGN.CENTER), ("Viscosity against temperature", False, INK, PP_ALIGN.LEFT)],
             [("3", True, NAVY, PP_ALIGN.CENTER), ("Viscosity against pressure", False, INK, PP_ALIGN.LEFT)]],
            [Inches(0.65), Inches(5.30)], head=False, row_h=Inches(0.36), size=12.5)
tb, tf = textbox(s, R, end + Inches(0.12), COLW, Inches(0.5))
para(tf, "Kinematic viscosity is the figure normally quoted; ISO VG is the grading standard.",
     12.5, False, GREY, 0, first=True, line=1.2)

y3 = end + Inches(0.75)
y3 = chip(s, R, y3, COLW, "Grease = base oil + thickener + additives")
tb, tf = textbox(s, R, y3, COLW, Inches(1.6))
para(tf, "Think of a sponge holding water. Squeeze it and water comes out; release it "
         "and the sponge takes the water back.", 13, False, INK, 8, first=True, line=1.25)
para(tf, "The thickener matrix holds the base oil until the bearing turns and load comes "
         "on. The oil is pushed out, and it is the base oil that does the lubricating.",
     13, False, INK, 0, line=1.25)

callout(s, R, y3 + Inches(1.75), COLW,
        "Consistency is graded by worked penetration (NLGI). NLGI 2 is the general "
        "multipurpose grade.", h=Inches(0.8), size=13, bold=True, color=NAVY)

prs.save(OUT)
print("saved", OUT)
