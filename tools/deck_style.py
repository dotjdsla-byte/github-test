# -*- coding: utf-8 -*-
"""Shared slide furniture for the KLC-SM house style.

A navy title band across the top, a plain white body laid out on a two-column
grid, grey chips labelling each block, and the internal-classification footer.
"""
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

NAVY = RGBColor(0x17, 0x2A, 0x4A)
INK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x59, 0x59, 0x59)
LIGHT = RGBColor(0xEF, 0xEF, 0xEF)
LINE = RGBColor(0xD0, 0xD4, 0xDA)
ACCENT = RGBColor(0x1F, 0x6F, 0x3D)
RED = RGBColor(0xC0, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

EA_FONT = "맑은 고딕"     # East Asian face, used for Korean runs
LATIN_FONT = "맑은 고딕"  # set to "Calibri" for an English-only deck

SW, SH = Inches(13.333), Inches(7.5)
BAND_H = Inches(1.12)
L = Inches(0.55)          # left column x
COLW = Inches(5.95)       # column width
R = Inches(6.83)          # right column x
TOP = Inches(1.45)        # first content row
FOOT = ("Classified as Internal - This content is proprietary information "
        "intended for employees and partners only.")


def kfont(run, size, bold=False, color=INK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = LATIN_FONT
    rPr = run._r.get_or_add_rPr()
    for tag, face in (("a:ea", EA_FONT), ("a:cs", LATIN_FONT)):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", face)


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = 0
    return tb, tf


def para(tf, text, size, bold=False, color=INK, space_after=4, first=False,
         align=PP_ALIGN.LEFT, line=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    if line:
        p.line_spacing = line
    r = p.add_run()
    r.text = text
    kfont(r, size, bold, color)
    return p


def rect(slide, x, y, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    return sh


def chip(slide, x, y, w, text):
    """Light grey section label; returns the y to start content at."""
    h = Inches(0.36)
    rect(slide, x, y, w, h, LIGHT)
    tb, tf = textbox(slide, x + Inches(0.12), y, w - Inches(0.2), h,
                     anchor=MSO_ANCHOR.MIDDLE)
    para(tf, text, 12, True, RGBColor(0x33, 0x33, 0x33), 0, first=True)
    return y + h + Inches(0.14)


def frame(slide, title):
    rect(slide, 0, 0, SW, BAND_H, NAVY)
    tb, tf = textbox(slide, Inches(0.55), Inches(0.16), Inches(12.2),
                     BAND_H - Inches(0.3), anchor=MSO_ANCHOR.MIDDLE)
    para(tf, title, 25, True, WHITE, 0, first=True)
    fb, ff = textbox(slide, Inches(0.55), SH - Inches(0.42), Inches(12.2), Inches(0.3))
    para(ff, FOOT, 10, False, RGBColor(0x8C, 0x8C, 0x8C), 0, first=True)


def table(slide, x, y, w, rows, col_w, head=True, row_h=Inches(0.34), size=11.5):
    """Bordered table drawn from rectangles; rows are str or (text, bold, colour, align)."""
    cy = y
    for ri, row in enumerate(rows):
        cx = x
        for ci, cell in enumerate(row):
            cw = col_w[ci]
            is_head = head and ri == 0
            rect(slide, cx, cy, cw, row_h,
                 NAVY if is_head else WHITE, None if is_head else LINE)
            tb, tf = textbox(slide, cx + Inches(0.08), cy, cw - Inches(0.16), row_h,
                             anchor=MSO_ANCHOR.MIDDLE)
            txt, bold, col, align = (cell if isinstance(cell, tuple)
                                     else (cell, False, INK, PP_ALIGN.LEFT))
            if is_head:
                bold, col = True, WHITE
            para(tf, txt, size, bold, col, 0, first=True, align=align)
            cx += cw
        cy += row_h
    return cy


def callout(slide, x, y, w, text, h=Inches(0.85), size=12.5, bold=False, color=INK):
    rect(slide, x, y, w, h, LIGHT)
    tb, tf = textbox(slide, x + Inches(0.2), y, w - Inches(0.4), h,
                     anchor=MSO_ANCHOR.MIDDLE)
    para(tf, text, size, bold, color, 0, first=True, line=1.25)
    return y + h


def blank_slide(prs, title):
    """Add a slide stripped of inherited placeholders, with the band and footer."""
    s = prs.slides.add_slide(prs.slide_layouts[6] if len(prs.slide_layouts) > 6
                             else prs.slide_layouts[0])
    for shp in list(s.shapes):
        shp._element.getparent().remove(shp._element)
    frame(s, title)
    return s


def notes(slide, text):
    """Speaker notes, written as plain spoken English."""
    slide.notes_slide.notes_text_frame.text = text
